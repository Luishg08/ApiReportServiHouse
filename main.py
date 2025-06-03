import webbrowser 
from fastapi import FastAPI, Depends, HTTPException
from sqlalchemy.orm import Session
from database import SessionLocal, engine
from models import Base, Order, Stock, StockTransaction, Delivery, Location
from fastapi.responses import StreamingResponse
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from datetime import datetime, date
from fastapi.middleware.cors import CORSMiddleware

#No se requieren crear las tablas.

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Cambia esto por ["http://localhost:3000"] en producción
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


# Prueba para verificar que FastAPI está funcionando bien
@app.get("/TestApiReport")
def read_root():
    return {"message": "Hola desde ApiReport"}

def get_delivered_orders_query(db, idDelivery=None, today_only=False):
    query = db.query(Order, Delivery).join(Delivery, Order.delivery_id == Delivery.id).filter(Order.state == "DELIVERED")
    if idDelivery:
        query = query.filter(Order.delivery_id == idDelivery)
    if today_only:
        today = date.today()
        query = query.filter(Order.updatedAt >= datetime.combine(today, datetime.min.time()),
                             Order.updatedAt <= datetime.combine(today, datetime.max.time()))
    return query.all()

def format_excel(ws):
    # Encabezado en negrita, fondo azul y letras blancas
    header_font = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4F81BD")
    align = Alignment(horizontal="center", vertical="center")
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = fill
        cell.alignment = align
        cell.border = border
    ws.freeze_panes = "A2"
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 18

def write_pdf_header(p, y, title):
    p.setFont("Helvetica-Bold", 16)
    p.setFillColor(colors.HexColor("#4F81BD"))
    p.drawString(50, y, title)
    p.setFillColor(colors.black)
    p.setFont("Helvetica-Bold", 10)
    y -= 25
    p.drawString(50, y, "Order ID")
    p.drawString(120, y, "Order Number")
    p.drawString(230, y, "Delivered At")
    p.drawString(340, y, "Delivery Name")
    p.drawString(470, y, "Delivery User ID")
    return y - 10

@app.get("/api/report/delivery/{idDelivery}/excel")
def report_delivery_excel(idDelivery: int, db: Session = Depends(get_db)):
    orders = get_delivered_orders_query(db, idDelivery=idDelivery, today_only=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Delivered Orders"
    ws.append([
        "Order Number", "Delivered At", "Delivery Name", "Delivery Email",
        "Recipient Email", "Product", "Quantity", "Delivery Address"
    ])
    format_excel(ws)
    for order, delivery in orders:
        recipient = getattr(order, "recipient", None)  # Ajusta si tienes destinatario
        recipient_email = getattr(order, "email", "")
        address = getattr(order.final_address, "address", "") if order.final_address else ""
        products = order.stockTransactions if hasattr(order, "stockTransactions") else []
        first = True
        for st in products:
            product_name = getattr(st.stock.product, "name", "") if st.stock and st.stock.product else ""
            quantity = getattr(st, "amount", "")
            ws.append([
                order.order_number if first else "",
                order.updatedAt.strftime("%Y-%m-%d %H:%M") if first else "",
                getattr(delivery, "full_name", "") if first else "",
                getattr(delivery, "email", "") if first else "",
                recipient_email if first else "",
                product_name,
                quantity,
                address if first else ""
            ])
            first = False
        if not products:
            ws.append([
                order.order_number,
                order.updatedAt.strftime("%Y-%m-%d %H:%M"),
                getattr(delivery, "full_name", ""),
                getattr(delivery, "email", ""),
                recipient_email,
                "",
                "",
                address
            ])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=delivery_{idDelivery}_today.xlsx"}
    )

@app.get("/api/report/delivery/{idDelivery}/pdf")
def report_delivery_pdf(idDelivery: int, db: Session = Depends(get_db)):
    orders = get_delivered_orders_query(db, idDelivery=idDelivery, today_only=True)
    buffer = io.BytesIO()
    p = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter
    y = height - 50
    p.setFont("Helvetica-Bold", 14)
    p.setFillColor(colors.HexColor("#4F81BD"))
    p.drawString(50, y, f"Delivered Orders - Delivery {idDelivery}")
    p.setFillColor(colors.black)
    y -= 30
    p.setFont("Helvetica-Bold", 10)
    for order, delivery in orders:
        if y < 120:
            p.showPage()
            y = height - 50
        address = getattr(order.final_address, "address", "") if order.final_address else ""
        p.drawString(50, y, f"Order Number: {order.order_number}")
        y -= 15
        p.drawString(50, y, f"Delivered At: {order.updatedAt.strftime('%Y-%m-%d %H:%M')}")
        y -= 15
        p.drawString(50, y, f"Delivery: {getattr(delivery, 'full_name', '')} | Email: {getattr(delivery, 'email', '')}")
        y -= 15
        p.drawString(50, y, f"Recipient Email: {order.email}")
        y -= 15
        p.drawString(50, y, f"Delivery Address: {address}")
        y -= 15
        p.setFont("Helvetica-Bold", 9)
        p.drawString(70, y, "Products:")
        y -= 12
        p.setFont("Helvetica", 9)
        products = order.stockTransactions if hasattr(order, "stockTransactions") else []
        for st in products:
            if y < 80:
                p.showPage()
                y = height - 50
            product_name = getattr(st.stock.product, "name", "") if st.stock and st.stock.product else ""
            quantity = getattr(st, "amount", "")
            p.drawString(90, y, f"- {product_name} (Cantidad: {quantity})")
            y -= 12
        y -= 10
        p.setStrokeColor(colors.HexColor("#4F81BD"))
        p.line(50, y, width - 50, y)
        y -= 20
        p.setFont("Helvetica-Bold", 10)
    p.save()
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=delivery_{idDelivery}_today.pdf"}
    )


@app.get("/api/report/all/excel")
def report_all_excel(db: Session = Depends(get_db)):
    orders = get_delivered_orders_query(db)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Delivered Orders"
    ws.append([
        "Order Number", "Delivered At", "Delivery Name", "Delivery Email",
        "Recipient Email", "Product", "Quantity", "Delivery Address"
    ])
    format_excel(ws)
    for order, delivery in orders:
        recipient_email = getattr(order, "email", "")
        address = getattr(order.final_address, "address", "") if order.final_address else ""
        products = order.stockTransactions if hasattr(order, "stockTransactions") else []
        first = True
        for st in products:
            # Acceso correcto al nombre del producto y cantidad
            product_name = getattr(st.stock.product, "name", "") if st.stock and st.stock.product else ""
            quantity = getattr(st, "amount", "")
            ws.append([
                order.order_number if first else "",
                order.updatedAt.strftime("%Y-%m-%d %H:%M") if first else "",
                getattr(delivery, "full_name", "") if first else "",
                getattr(delivery, "email", "") if first else "",
                recipient_email if first else "",
                product_name,
                quantity,
                address if first else ""
            ])
            first = False
        if not products:
            ws.append([
                order.order_number,
                order.updatedAt.strftime("%Y-%m-%d %H:%M"),
                getattr(delivery, "full_name", ""),
                getattr(delivery, "email", ""),
                recipient_email,
                "",
                "",
                address
            ])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=all_delivered_orders.xlsx"}
    )

@app.get("/api/report/all/pdf")
def report_all_pdf(db: Session = Depends(get_db)):
    orders = get_delivered_orders_query(db)
    buffer = io.BytesIO()
    p = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter
    y = height - 50
    p.setFont("Helvetica-Bold", 14)
    p.setFillColor(colors.HexColor("#4F81BD"))
    p.drawString(50, y, "Delivered Orders")
    p.setFillColor(colors.black)
    y -= 30
    p.setFont("Helvetica-Bold", 10)
    for order, delivery in orders:
        if y < 120:
            p.showPage()
            y = height - 50
        address = getattr(order.final_address, "address", "") if order.final_address else ""
        p.drawString(50, y, f"Order Number: {order.order_number}")
        y -= 15
        p.drawString(50, y, f"Delivered At: {order.updatedAt.strftime('%Y-%m-%d %H:%M')}")
        y -= 15
        p.drawString(50, y, f"Delivery: {getattr(delivery, 'full_name', '')} | Email: {getattr(delivery, 'email', '')}")
        y -= 15
        p.drawString(50, y, f"Recipient Email: {order.email}")
        y -= 15
        p.drawString(50, y, f"Delivery Address: {address}")
        y -= 15
        p.setFont("Helvetica-Bold", 9)
        p.drawString(70, y, "Products:")
        y -= 12
        p.setFont("Helvetica", 9)
        stocksTransactions = order.stockTransactions if hasattr(order, "stockTransactions") else []
        for st in stocksTransactions:
            if y < 80:
                p.showPage()
                y = height - 50
            product_name = getattr(st.stock.product, "name", "") if st.stock and st.stock.product else ""
            quantity = getattr(st, "amount", "")
            p.drawString(90, y, f"- {product_name} (Cantidad: {quantity})")
            y -= 12
        y -= 10
        p.setStrokeColor(colors.HexColor("#4F81BD"))
        p.line(50, y, width - 50, y)
        y -= 20
        p.setFont("Helvetica-Bold", 10)
    p.save()
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=all_delivered_orders.pdf"}
    )



if __name__ == "__main__":
    webbrowser.open("http://127.0.0.1:8088/docs") # Se abre automáticamente Swagger UI para la documentación de ApiReport

# Ejecutar con siguiente comando: $ uvicorn main:app --reload --port 8088 --host 0.0.0.0